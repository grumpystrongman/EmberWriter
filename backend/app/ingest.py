from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from ebooklib import ITEM_DOCUMENT, epub
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from . import binder
from .binder_models import BinderNodeCreate
from .revisions import record_revision
from .storage import save_text

SUPPORTED_UPLOADS = {".md", ".txt", ".docx", ".pdf", ".epub", ".rtf", ".html", ".htm"}
CHAPTER_RE = re.compile(
    r"^(?:chapter\s+(?:\d+|[ivxlcdm]+|one|two|three|four|five|six|seven|eight|nine|ten)|"
    r"prologue|epilogue|interlude|part\s+(?:\d+|[ivxlcdm]+))\b.*$",
    re.IGNORECASE,
)


def _runs_to_markdown(paragraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        if run.bold and run.italic:
            text = f"***{text}***"
        elif run.bold:
            text = f"**{text}**"
        elif run.italic:
            text = f"*{text}*"
        parts.append(text)
    return "".join(parts).strip()


def _docx_to_markdown(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = _runs_to_markdown(paragraph)
        if not text:
            lines.append("")
            continue
        style = (paragraph.style.name or "").casefold() if paragraph.style else ""
        heading = re.match(r"heading\s+(\d+)", style)
        if heading:
            level = max(1, min(6, int(heading.group(1))))
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            lines.append(" | ".join(cells))
    return "\n\n".join(line for line in lines if line is not None).strip()


def _pdf_to_markdown(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    return "\n\n".join(pages).strip()


def _html_to_markdown(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav"]):
        tag.decompose()
    blocks: list[str] = []
    root = soup.body or soup
    for element in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "blockquote", "li"]):
        text = element.get_text(" ", strip=True)
        if not text:
            continue
        if element.name and element.name.startswith("h"):
            blocks.append(f"{'#' * int(element.name[1])} {text}")
        elif element.name == "blockquote":
            blocks.append(f"> {text}")
        elif element.name == "li":
            blocks.append(f"- {text}")
        else:
            blocks.append(text)
    return "\n\n".join(blocks).strip()


def _epub_to_markdown(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".epub") as handle:
        handle.write(data)
        handle.flush()
        book = epub.read_epub(handle.name)
    sections: list[str] = []
    for item in book.get_items_of_type(ITEM_DOCUMENT):
        converted = _html_to_markdown(item.get_content().decode("utf-8", errors="ignore"))
        if converted:
            sections.append(converted)
    return "\n\n".join(sections).strip()


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.casefold()
    if suffix not in SUPPORTED_UPLOADS:
        raise ValueError(f"Unsupported upload type: {suffix or 'unknown'}")
    if suffix == ".docx":
        return _docx_to_markdown(data)
    if suffix == ".pdf":
        return _pdf_to_markdown(data)
    if suffix == ".epub":
        return _epub_to_markdown(data)
    decoded = data.decode("utf-8-sig", errors="replace")
    if suffix == ".rtf":
        return rtf_to_text(decoded).strip()
    if suffix in {".html", ".htm"}:
        return _html_to_markdown(decoded)
    return decoded.strip()


def _chapter_title(line: str, fallback: str) -> str:
    cleaned = re.sub(r"^#+\s*", "", line).strip()
    return cleaned[:160] or fallback


def split_novel(markdown: str, fallback_title: str) -> list[tuple[str, str]]:
    lines = markdown.splitlines()
    starts: list[tuple[int, str]] = []
    for index, line in enumerate(lines):
        stripped = line.strip()
        plain = re.sub(r"^#+\s*", "", stripped)
        if stripped.startswith("# ") or CHAPTER_RE.match(plain):
            starts.append((index, _chapter_title(stripped, f"Section {len(starts) + 1}")))

    if len(starts) < 2:
        return [(fallback_title, markdown.strip())]

    parts: list[tuple[str, str]] = []
    if starts[0][0] > 0:
        preface = "\n".join(lines[: starts[0][0]]).strip()
        if preface:
            parts.append(("Front Matter", preface))
    for position, (start, title) in enumerate(starts):
        end = starts[position + 1][0] if position + 1 < len(starts) else len(lines)
        content = "\n".join(lines[start:end]).strip()
        if content:
            parts.append((title, content))
    return parts


def _root_id(state, title: str) -> str:
    node_map = {node.id: node for node in state.nodes}
    for node_id in state.roots:
        node = node_map.get(node_id)
        if node and node.title == title:
            return node.id
    raise ValueError(f"Binder root not found: {title}")


def _create_imported_node(slug: str, title: str, content: str, mode: str) -> dict:
    state = binder.get_binder(slug)
    if mode in {"novel", "portion"}:
        kind = "document"
        parent_id = _root_id(state, "Draft")
        include = True
    elif mode == "research":
        kind = "research"
        parent_id = _root_id(state, "Research")
        include = False
    else:
        kind = "note"
        parent_id = _root_id(state, "Story Bible")
        include = False

    before = {node.id for node in state.nodes}
    state = binder.create_node(
        slug,
        BinderNodeCreate(
            title=title,
            kind=kind,
            parent_id=parent_id,
            include_in_compile=include,
        ),
    )
    created = next(node for node in state.nodes if node.id not in before)
    if not created.path:
        raise ValueError("Imported Binder document did not receive a source path")
    save_text(slug, created.path, content.rstrip() + "\n")
    revision = record_revision(
        slug,
        created.path,
        content.rstrip() + "\n",
        source="import",
        note=f"Imported as {mode}",
        force=True,
    )
    return {
        "node_id": created.id,
        "title": created.title,
        "path": created.path,
        "word_count": len(content.split()),
        "revision_id": revision["id"],
    }


def import_bytes(slug: str, filename: str, data: bytes, mode: str = "novel") -> dict:
    if mode not in {"novel", "portion", "idea", "research"}:
        raise ValueError("Import mode must be novel, portion, idea, or research")
    text = extract_text(filename, data)
    if not text.strip():
        raise ValueError("The uploaded file did not contain readable text")
    fallback = Path(filename).stem.replace("-", " ").replace("_", " ").strip().title() or "Imported Text"
    parts = split_novel(text, fallback) if mode == "novel" else [(fallback, text)]
    imported = [_create_imported_node(slug, title, content, mode) for title, content in parts]
    return {
        "filename": filename,
        "mode": mode,
        "documents": imported,
        "total_words": sum(item["word_count"] for item in imported),
    }


def import_pasted_text(slug: str, title: str, content: str, mode: str = "idea") -> dict:
    if not content.strip():
        raise ValueError("Content is empty")
    if mode == "novel":
        parts = split_novel(content, title or "Imported Novel")
        imported = [_create_imported_node(slug, part_title, body, mode) for part_title, body in parts]
    else:
        imported = [_create_imported_node(slug, title or "Imported Text", content, mode)]
    return {
        "filename": None,
        "mode": mode,
        "documents": imported,
        "total_words": sum(item["word_count"] for item in imported),
    }
