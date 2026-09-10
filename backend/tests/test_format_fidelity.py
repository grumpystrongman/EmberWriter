from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from ebooklib import ITEM_DOCUMENT, epub
from pypdf import PdfReader

from app import binder, ingest, publishing, storage
from app.binder_models import BinderNodeUpdate


def use_temp_data(tmp_path: Path) -> None:
    storage.DATA_ROOT = tmp_path
    storage.PROJECTS_ROOT = tmp_path / "projects"


def test_docx_import_preserves_supported_inline_formatting_and_alignment() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    bold = paragraph.add_run("Bold")
    bold.bold = True
    paragraph.add_run(" ")
    italic = paragraph.add_run("Italic")
    italic.italic = True
    paragraph.add_run(" ")
    under = paragraph.add_run("Under")
    under.underline = True
    paragraph.add_run(" ")
    strike = paragraph.add_run("Strike")
    strike.font.strike = True
    paragraph.add_run(" ")
    marked = paragraph.add_run("Marked")
    marked.font.highlight_color = WD_COLOR_INDEX.YELLOW

    buffer = io.BytesIO()
    document.save(buffer)
    converted = ingest.extract_text("formatted.docx", buffer.getvalue())

    assert '<p style="text-align: center">' in converted
    assert "<strong>Bold</strong>" in converted
    assert "<em>Italic</em>" in converted
    assert "<u>Under</u>" in converted
    assert "<s>Strike</s>" in converted
    assert "<mark>Marked</mark>" in converted


def test_novel_split_prefers_real_chapter_markers_over_book_title_heading() -> None:
    markdown = """# The Book Title

A short dedication.

# Chapter 1

First chapter.

# Chapter 2

Second chapter.
"""
    parts = ingest.split_novel(markdown, "Fallback")
    assert [title for title, _ in parts] == ["Front Matter", "Chapter 1", "Chapter 2"]
    assert "The Book Title" in parts[0][1]


def test_publish_preserves_supported_rich_formatting_and_metadata(tmp_path: Path) -> None:
    use_temp_data(tmp_path)
    project = storage.create_project("Format Fidelity")
    slug = project["slug"]
    state = binder.get_binder(slug)
    first = next(node for node in state.nodes if node.path == "manuscript/chapter-001.md")
    assert first.path

    source = """# Chapter One

**Bold** *Italic* <u>Under</u> <s>Strike</s> <mark>Marked</mark>

<p style="text-align: center">Centered line</p>

> Quoted line

- First item
- Second item

---

After the break.
"""
    storage.save_text(slug, first.path, source)
    binder.update_node(slug, first.id, BinderNodeUpdate(title="Chapter One"))

    result = publishing.export_project(
        slug,
        formats=["docx", "epub", "pdf"],
        title="The Fidelity Test",
        author="A. Writer",
        language="en",
        trim_width=6,
        trim_height=9,
    )
    root = storage.project_root(slug)
    paths = {
        artifact["format"]: root / artifact["relative_path"]
        for artifact in result["artifacts"]
    }

    exported_docx = Document(paths["docx"])
    assert exported_docx.core_properties.title == "The Fidelity Test"
    assert exported_docx.core_properties.author == "A. Writer"
    runs = {
        run.text.strip(): run
        for paragraph in exported_docx.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    }
    assert runs["Bold"].bold is True
    assert runs["Italic"].italic is True
    assert runs["Under"].underline is True
    assert runs["Strike"].font.strike is True
    assert runs["Marked"].font.highlight_color == WD_COLOR_INDEX.YELLOW
    centered = next(paragraph for paragraph in exported_docx.paragraphs if paragraph.text == "Centered line")
    assert centered.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert any(paragraph.text == "* * *" for paragraph in exported_docx.paragraphs)

    exported_epub = epub.read_epub(str(paths["epub"]))
    epub_html = "\n".join(
        item.get_content().decode("utf-8", errors="ignore")
        for item in exported_epub.get_items_of_type(ITEM_DOCUMENT)
    )
    assert "<strong>Bold</strong>" in epub_html
    assert "<em>Italic</em>" in epub_html
    assert "<u>Under</u>" in epub_html
    assert "Centered line" in epub_html

    pdf = PdfReader(str(paths["pdf"]))
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    for expected in ("Bold", "Italic", "Under", "Strike", "Marked", "Centered line", "Quoted line"):
        assert expected in pdf_text
    first_page = pdf.pages[0]
    assert abs(float(first_page.mediabox.width) - 432.0) < 1
    assert abs(float(first_page.mediabox.height) - 648.0) < 1
