from __future__ import annotations

import html
import re
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import markdown as markdown_lib
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches, Pt
from ebooklib import epub
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from .binder import get_binder
from .storage import project_root, read_text


def _slug(value: str) -> str:
    result = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip().lower()).strip("-")
    return result or "manuscript"


def _children(state, parent_id: str):
    return sorted(
        (node for node in state.nodes if node.parent_id == parent_id),
        key=lambda node: node.position,
    )


def compiled_documents(slug: str) -> list[dict]:
    state = get_binder(slug)
    node_map = {node.id: node for node in state.nodes}
    draft = next(
        (node_map[node_id] for node_id in state.roots if node_map[node_id].title == "Draft"),
        None,
    )
    if draft is None:
        raise ValueError("Binder Draft root is missing")

    result: list[dict] = []

    def visit(parent_id: str) -> None:
        for node in _children(state, parent_id):
            if node.kind == "folder":
                visit(node.id)
                continue
            if (
                not node.include_in_compile
                or not node.path
                or node.custom_metadata.get("source_missing")
            ):
                continue
            try:
                content = read_text(slug, node.path)
            except (FileNotFoundError, OSError, ValueError, UnicodeError):
                continue
            result.append(
                {
                    "id": node.id,
                    "title": node.title,
                    "path": node.path,
                    "content": content,
                    "word_count": node.word_count,
                    "metadata": node.custom_metadata,
                }
            )

    visit(draft.id)
    if not result:
        raise ValueError("No Binder documents are enabled for Compile")
    return result


def _markdown_html(content: str) -> str:
    return markdown_lib.markdown(content, extensions=["extra", "sane_lists"])


def _clean_chapter_content(content: str, title: str) -> str:
    lines = content.splitlines()
    if lines and re.match(r"^#\s+", lines[0].strip()):
        heading = re.sub(r"^#\s+", "", lines[0].strip()).strip()
        if heading.casefold() == title.strip().casefold():
            return "\n".join(lines[1:]).lstrip()
    return content


def _alignment_from_style(element: Tag) -> str | None:
    style = str(element.get("style", ""))
    match = re.search(r"text-align\s*:\s*(left|center|right|justify)", style, re.IGNORECASE)
    return match.group(1).casefold() if match else None


def _docx_alignment(element: Tag):
    alignment = _alignment_from_style(element)
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(alignment)


def _append_docx_inline(
    paragraph,
    node,
    *,
    bold: bool = False,
    italic: bool = False,
    underline: bool = False,
    strike: bool = False,
    highlight: bool = False,
) -> None:
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.underline = underline
        run.font.strike = strike
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    if not isinstance(node, Tag):
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return

    name = node.name.casefold() if node.name else ""
    child_bold = bold or name in {"b", "strong"}
    child_italic = italic or name in {"i", "em"}
    child_underline = underline or name == "u"
    child_strike = strike or name in {"s", "strike", "del"}
    child_highlight = highlight or name == "mark"
    for child in node.children:
        _append_docx_inline(
            paragraph,
            child,
            bold=child_bold,
            italic=child_italic,
            underline=child_underline,
            strike=child_strike,
            highlight=child_highlight,
        )


def _docx_paragraph_from_element(doc: Document, element: Tag, *, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    alignment = _docx_alignment(element)
    if alignment is not None:
        paragraph.alignment = alignment
    for child in element.children:
        _append_docx_inline(paragraph, child)
    return paragraph


def _write_docx_block(doc: Document, element: Tag) -> None:
    name = element.name.casefold() if element.name else ""
    if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        level = min(3, max(1, int(name[1])))
        paragraph = doc.add_paragraph(style=f"Heading {level}")
        alignment = _docx_alignment(element)
        if alignment is not None:
            paragraph.alignment = alignment
        for child in element.children:
            _append_docx_inline(paragraph, child)
        return
    if name in {"ul", "ol"}:
        list_style = "List Bullet" if name == "ul" else "List Number"
        for item in element.find_all("li", recursive=False):
            _docx_paragraph_from_element(doc, item, style=list_style)
        return
    if name == "blockquote":
        paragraph = _docx_paragraph_from_element(doc, element)
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.right_indent = Inches(0.2)
        return
    if name == "hr":
        paragraph = doc.add_paragraph("* * *")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    if name == "pre":
        paragraph = _docx_paragraph_from_element(doc, element)
        for run in paragraph.runs:
            run.font.name = "Courier New"
        return
    _docx_paragraph_from_element(doc, element)


def export_docx(path: Path, *, title: str, author: str, documents: list[dict]) -> None:
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = author
    doc.core_properties.subject = "Compiled by EmberWriter"

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(22)
    if author:
        author_para = doc.add_paragraph()
        author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_run = author_para.add_run(author)
        author_run.italic = True
    doc.add_page_break()

    for index, item in enumerate(documents):
        if index:
            doc.add_page_break()
        heading = doc.add_heading(item["title"], level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        body = _clean_chapter_content(item["content"], item["title"])
        soup = BeautifulSoup(_markdown_html(body), "html.parser")
        for element in soup.find_all(recursive=False):
            if isinstance(element, Tag):
                _write_docx_block(doc, element)
    doc.save(path)


def export_epub(
    path: Path,
    *,
    title: str,
    author: str,
    language: str,
    documents: list[dict],
    include_toc: bool,
) -> None:
    book = epub.EpubBook()
    book.set_identifier(f"urn:uuid:{uuid4()}")
    book.set_title(title)
    book.set_language(language or "en")
    if author:
        book.add_author(author)

    chapters = []
    for index, item in enumerate(documents, start=1):
        filename = f"chapter-{index:04d}.xhtml"
        chapter = epub.EpubHtml(title=item["title"], file_name=filename, lang=language or "en")
        body = _clean_chapter_content(item["content"], item["title"])
        chapter.content = (
            f"<h1>{html.escape(item['title'])}</h1>\n" + _markdown_html(body)
        ).encode("utf-8")
        book.add_item(chapter)
        chapters.append(chapter)

    book.toc = tuple(chapters) if include_toc else ()
    book.spine = ["nav", *chapters]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    style = (
        "body { font-family: serif; line-height: 1.45; } "
        "h1 { text-align: center; margin: 2em 0; } "
        "blockquote { margin-left: 1.5em; margin-right: 1.5em; }"
    )
    css = epub.EpubItem(
        uid="style",
        file_name="style/book.css",
        media_type="text/css",
        content=style,
    )
    book.add_item(css)
    for chapter in chapters:
        chapter.add_item(css)
    epub.write_epub(str(path), book)


def _reportlab_inline(node) -> str:
    if isinstance(node, NavigableString):
        return html.escape(str(node))
    if not isinstance(node, Tag):
        return ""
    if node.name == "br":
        return "<br/>"
    content = "".join(_reportlab_inline(child) for child in node.children)
    name = node.name.casefold() if node.name else ""
    if name in {"b", "strong"}:
        return f"<b>{content}</b>"
    if name in {"i", "em"}:
        return f"<i>{content}</i>"
    if name == "u":
        return f"<u>{content}</u>"
    if name in {"s", "strike", "del"}:
        return f"<strike>{content}</strike>"
    if name == "code":
        return f'<font name="Courier">{content}</font>'
    # Highlight is an editorial aid; preserve the text but do not print the highlight background.
    return content


def _pdf_alignment(element: Tag) -> int:
    return {
        "left": TA_LEFT,
        "center": TA_CENTER,
        "right": TA_RIGHT,
    }.get(_alignment_from_style(element), TA_LEFT)


def _pdf_page_number(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 9)
    canvas.drawCentredString(doc.pagesize[0] / 2, 0.35 * inch, str(doc.page))
    canvas.restoreState()


def export_pdf(
    path: Path,
    *,
    title: str,
    author: str,
    documents: list[dict],
    trim_width: float,
    trim_height: float,
) -> None:
    page_size = (trim_width * inch, trim_height * inch)
    doc = SimpleDocTemplate(
        str(path),
        pagesize=page_size,
        rightMargin=0.7 * inch,
        leftMargin=0.9 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=title,
        author=author,
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BookBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=10.5,
        leading=14,
        spaceAfter=7,
        firstLineIndent=16,
        allowWidows=0,
        allowOrphans=0,
    )
    quote_style = ParagraphStyle(
        "BookQuote",
        parent=body,
        leftIndent=20,
        rightIndent=12,
        firstLineIndent=0,
    )
    chapter_style = ParagraphStyle(
        "ChapterTitle",
        parent=styles["Heading1"],
        fontName="Times-Bold",
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        spaceAfter=28,
    )
    title_style = ParagraphStyle(
        "BookTitle",
        parent=chapter_style,
        fontSize=24,
        leading=30,
        spaceBefore=1.8 * inch,
    )
    story: list[object] = [Paragraph(html.escape(title), title_style)]
    if author:
        story.extend([Spacer(1, 0.25 * inch), Paragraph(html.escape(author), chapter_style)])
    story.append(PageBreak())

    for index, item in enumerate(documents):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(html.escape(item["title"]), chapter_style))
        body_markdown = _clean_chapter_content(item["content"], item["title"])
        soup = BeautifulSoup(_markdown_html(body_markdown), "html.parser")
        for element in soup.find_all(recursive=False):
            if not isinstance(element, Tag):
                continue
            name = element.name.casefold() if element.name else ""
            if name == "hr":
                break_style = ParagraphStyle("SceneBreak", parent=body, alignment=TA_CENTER, firstLineIndent=0)
                story.append(Paragraph("* * *", break_style))
                continue
            if name in {"ul", "ol"}:
                items = []
                for child in element.find_all("li", recursive=False):
                    items.append(ListItem(Paragraph(_reportlab_inline(child), body)))
                if items:
                    story.append(
                        ListFlowable(
                            items,
                            bulletType="bullet" if name == "ul" else "1",
                            leftIndent=24,
                        )
                    )
                continue
            if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                heading_style = ParagraphStyle(
                    f"BodyHeading{name[1]}",
                    parent=styles["Heading2"],
                    alignment=_pdf_alignment(element),
                    spaceBefore=10,
                    spaceAfter=6,
                )
                story.append(Paragraph(_reportlab_inline(element), heading_style))
                continue
            paragraph_style = quote_style if name == "blockquote" else body
            alignment = _pdf_alignment(element)
            if alignment != paragraph_style.alignment:
                paragraph_style = ParagraphStyle(
                    f"Aligned-{alignment}-{len(story)}",
                    parent=paragraph_style,
                    alignment=alignment,
                    firstLineIndent=0 if alignment != TA_LEFT else paragraph_style.firstLineIndent,
                )
            markup = _reportlab_inline(element).strip()
            if markup:
                story.append(Paragraph(markup, paragraph_style))

    doc.build(story, onLaterPages=_pdf_page_number)


def export_project(
    slug: str,
    *,
    formats: list[str],
    title: str,
    author: str = "",
    language: str = "en",
    include_toc: bool = True,
    trim_width: float = 6.0,
    trim_height: float = 9.0,
) -> dict:
    allowed = {"docx", "epub", "pdf"}
    normalized = []
    for item in formats:
        value = item.casefold().strip().lstrip(".")
        if value not in allowed:
            raise ValueError(f"Unsupported publish format: {item}")
        if value not in normalized:
            normalized.append(value)
    if not normalized:
        raise ValueError("Choose at least one publish format")
    if not 4.0 <= trim_width <= 8.5 or not 6.0 <= trim_height <= 11.7:
        raise ValueError("Print trim size is outside supported book dimensions")

    documents = compiled_documents(slug)
    export_id = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8]
    directory = project_root(slug) / "exports" / export_id
    directory.mkdir(parents=True, exist_ok=True)
    base = _slug(title)
    artifacts = []

    for fmt in normalized:
        path = directory / f"{base}.{fmt}"
        if fmt == "docx":
            export_docx(path, title=title, author=author, documents=documents)
        elif fmt == "epub":
            export_epub(
                path,
                title=title,
                author=author,
                language=language,
                documents=documents,
                include_toc=include_toc,
            )
        else:
            export_pdf(
                path,
                title=title,
                author=author,
                documents=documents,
                trim_width=trim_width,
                trim_height=trim_height,
            )
        artifacts.append(
            {
                "format": fmt,
                "filename": path.name,
                "relative_path": str(path.relative_to(project_root(slug))).replace("\\", "/"),
                "bytes": path.stat().st_size,
            }
        )

    return {
        "export_id": export_id,
        "title": title,
        "documents": len(documents),
        "words": sum(item["word_count"] for item in documents),
        "artifacts": artifacts,
    }
