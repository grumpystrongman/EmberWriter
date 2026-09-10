from __future__ import annotations

import hashlib
import html
import json
import re
import sqlite3
import zipfile
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

from bs4 import BeautifulSoup, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

from .distribution import load_release_profile
from .generation import generate
from .memory import chapter_order
from .models import ProviderConfig
from .publishing import (
    _append_docx_inline,
    _clean_chapter_content,
    _markdown_html,
    _reportlab_inline,
    compiled_documents,
)
from .storage import _connect, project_root, utc_now
from .submission_models import (
    SubmissionBuildResponse,
    SubmissionDestination,
    SubmissionPackageArtifact,
    SubmissionProfile,
    SubmissionRecord,
    SubmissionValidationIssue,
    SubmissionValidationResponse,
)

SHUNN_FORMAT_URL = "https://www.shunn.net/format/"
RULES_REVIEWED = "2026-09-10"


def _profile_path(slug: str) -> Path:
    return project_root(slug) / "publishing" / "submission-profile.json"


def _ensure_project(slug: str) -> Path:
    root = project_root(slug)
    if not (root / "project.json").exists():
        raise FileNotFoundError(slug)
    return root


def _word_count(documents: list[dict]) -> int:
    return sum(len(re.findall(r"\b[\w’'-]+\b", item["content"])) for item in documents)


def _compiled_or_empty(slug: str) -> list[dict]:
    try:
        return compiled_documents(slug)
    except ValueError as exc:
        if str(exc) != "No Binder documents are enabled for Compile":
            raise
        return []


def default_submission_profile(slug: str) -> SubmissionProfile:
    _ensure_project(slug)
    documents = _compiled_or_empty(slug)
    release = load_release_profile(slug)
    return SubmissionProfile(
        author_name=release.author,
        title=release.title,
        genre="",
        word_count=_word_count(documents),
        bio=release.author_bio,
        comp_titles=[],
    )


def load_submission_profile(slug: str) -> SubmissionProfile:
    path = _profile_path(slug)
    if not path.exists():
        return default_submission_profile(slug)
    return SubmissionProfile.model_validate_json(path.read_text(encoding="utf-8"))


def save_submission_profile(slug: str, profile: SubmissionProfile) -> SubmissionProfile:
    _ensure_project(slug)
    path = _profile_path(slug)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".json.tmp")
    temporary.write_text(
        json.dumps(profile.model_dump(), indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)
    return profile


def _destination(profile: SubmissionProfile, destination_id: str) -> SubmissionDestination:
    for destination in profile.destinations:
        if destination.id == destination_id:
            return destination
    raise FileNotFoundError(destination_id)


def _valid_iso_date(value: str) -> bool:
    if not value:
        return True
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def validate_submission(
    slug: str,
    profile: SubmissionProfile,
    destination_id: str,
) -> SubmissionValidationResponse:
    _ensure_project(slug)
    destination = _destination(profile, destination_id)
    issues: list[SubmissionValidationIssue] = []

    def add(level: str, code: str, message: str) -> None:
        issues.append(
            SubmissionValidationIssue(
                level=level,
                code=code,
                message=message,
                destination_id=destination.id,
            )
        )

    if not profile.title.strip():
        add("error", "title-required", "Enter the manuscript title.")
    if not profile.author_name.strip():
        add("error", "author-required", "Enter the author or pen name used for this submission.")
    if not profile.genre.strip():
        add("warning", "genre-missing", "Add the manuscript genre/category used in the query.")
    if profile.word_count <= 0:
        add("warning", "word-count-missing", "Confirm the manuscript word count before sending the query.")
    if destination.query_required and not profile.query_letter.strip():
        add("error", "query-required", f"{destination.name} requires a query letter.")
    if destination.synopsis_required and not profile.synopsis.strip():
        add("error", "synopsis-required", f"{destination.name} requires a synopsis.")
    if destination.bio_required and not profile.bio.strip():
        add("error", "bio-required", f"{destination.name} requires an author bio.")
    if destination.sample_kind != "none" and destination.sample_count <= 0:
        add("error", "sample-count", "The requested sample amount must be greater than zero.")
    if destination.sample_kind == "pages":
        add(
            "info",
            "page-boundary",
            "Page-count samples are cut from Ember's standardized 8.5×11 submission PDF. "
            "Word processors can repaginate DOCX files, so verify the destination's preferred "
            "page boundary before sending.",
        )
    if destination.method == "email" and not destination.email.strip():
        add(
            "warning",
            "email-missing",
            "This destination is marked as email submission but has no email address saved.",
        )
    if destination.method in {"query_manager", "web_form"} and not destination.submission_url.strip():
        add(
            "warning",
            "submission-url-missing",
            "Save the current submission portal URL for this destination.",
        )
    if not destination.guidelines_url.strip():
        add(
            "warning",
            "guidelines-source-missing",
            "Save the destination's current guidelines URL. Individual agent/publisher "
            "requirements override Ember's manuscript preset.",
        )
    if not destination.accepted_formats and destination.attachment_mode != "body":
        add("error", "formats-required", "Attachment submissions need at least one accepted file format.")
    unsupported = {
        value.casefold().lstrip(".")
        for value in destination.accepted_formats
        if value.casefold().lstrip(".") not in {"docx", "pdf", "txt"}
    }
    if unsupported:
        add(
            "warning",
            "unsupported-format",
            "Ember cannot generate these requested attachment formats yet: "
            + ", ".join(sorted(unsupported)),
        )
    for record in profile.records:
        for field_name, value in (
            ("submitted_at", record.submitted_at),
            ("follow_up_on", record.follow_up_on),
            ("response_at", record.response_at),
        ):
            if value and not _valid_iso_date(value[:10]):
                add(
                    "error",
                    "record-date",
                    f"Submission record {field_name} must begin with YYYY-MM-DD.",
                )
                break

    sample_description = {
        "none": "No manuscript sample",
        "pages": f"First {destination.sample_count} standardized manuscript pages",
        "chapters": f"First {destination.sample_count} Binder chapters/documents",
        "words": f"First {destination.sample_count:,} manuscript words",
        "full": "Full compiled manuscript",
    }[destination.sample_kind]
    return SubmissionValidationResponse(
        valid=not any(issue.level == "error" for issue in issues),
        issues=issues,
        destination_id=destination.id,
        sample_description=sample_description,
    )


def _font_for_preset(preset: str) -> str:
    return "Courier New" if preset == "shunn_classic" else "Times New Roman"


def _short_title(title: str) -> str:
    words = re.findall(r"[A-Za-z0-9]+", title.upper())
    return " ".join(words[:3]) or "MANUSCRIPT"


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def _configure_submission_doc(
    doc: Document,
    *,
    profile: SubmissionProfile,
    include_header: bool,
) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = _font_for_preset(profile.manuscript_preset)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)
    if include_header:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        surname = profile.author_name.strip().split()[-1] if profile.author_name.strip() else "AUTHOR"
        header.add_run(f"{surname} / {_short_title(profile.title)} / ")
        _add_page_field(header)


def _write_manuscript_docx(
    path: Path,
    *,
    profile: SubmissionProfile,
    documents: list[dict],
    include_title_page: bool,
) -> None:
    doc = Document()
    _configure_submission_doc(doc, profile=profile, include_header=True)
    if include_title_page:
        contact = doc.add_paragraph()
        contact.paragraph_format.first_line_indent = Inches(0)
        contact.paragraph_format.line_spacing = 1
        for line in [
            profile.author_name,
            profile.address,
            profile.phone,
            profile.email,
            profile.website,
        ]:
            if line.strip():
                contact.add_run(line.strip()).add_break()
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = Inches(0)
        title.paragraph_format.space_before = Inches(2)
        title.add_run(profile.title.upper()).bold = True
        if profile.genre or profile.word_count:
            meta = doc.add_paragraph()
            meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
            meta.paragraph_format.first_line_indent = Inches(0)
            pieces = [
                profile.genre.strip(),
                f"approximately {profile.word_count:,} words" if profile.word_count else "",
            ]
            meta.add_run(" · ".join(piece for piece in pieces if piece))
        doc.add_page_break()

    for doc_index, item in enumerate(documents):
        if doc_index:
            doc.add_page_break()
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.first_line_indent = Inches(0)
        heading.paragraph_format.line_spacing = 2
        heading.add_run(item["title"])
        body = _clean_chapter_content(item["content"], item["title"])
        soup = BeautifulSoup(_markdown_html(body), "html.parser")
        for element in soup.find_all(recursive=False):
            if not isinstance(element, Tag):
                continue
            if element.name == "hr":
                scene_break = doc.add_paragraph("#")
                scene_break.alignment = WD_ALIGN_PARAGRAPH.CENTER
                scene_break.paragraph_format.first_line_indent = Inches(0)
                scene_break.paragraph_format.line_spacing = 2
                continue
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.line_spacing = 2
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.first_line_indent = Inches(0.5)
            if element.name in {"blockquote", "pre"}:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.left_indent = Inches(0.5)
            _append_docx_inline(paragraph, element)
    doc.save(path)


def _pdf_page_header(canvas, document, *, profile: SubmissionProfile) -> None:
    canvas.saveState()
    font = "Courier" if profile.manuscript_preset == "shunn_classic" else "Times-Roman"
    canvas.setFont(font, 9)
    surname = profile.author_name.strip().split()[-1] if profile.author_name.strip() else "AUTHOR"
    label = f"{surname} / {_short_title(profile.title)} / {document.page}"
    canvas.drawRightString(LETTER[0] - inch, LETTER[1] - 0.55 * inch, label)
    canvas.restoreState()


def _write_manuscript_pdf(
    path: Path,
    *,
    profile: SubmissionProfile,
    documents: list[dict],
) -> None:
    font = "Courier" if profile.manuscript_preset == "shunn_classic" else "Times-Roman"
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "SubmissionBody",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=12,
        leading=24,
        firstLineIndent=0.5 * inch,
        spaceAfter=0,
        alignment=TA_LEFT,
    )
    heading = ParagraphStyle(
        "SubmissionHeading",
        parent=body,
        firstLineIndent=0,
        alignment=TA_CENTER,
        spaceAfter=24,
    )
    scene_break = ParagraphStyle(
        "SubmissionSceneBreak",
        parent=body,
        firstLineIndent=0,
        alignment=TA_CENTER,
    )
    flow: list[object] = []
    for doc_index, item in enumerate(documents):
        if doc_index:
            flow.append(PageBreak())
        flow.append(Paragraph(html.escape(item["title"]), heading))
        body_markdown = _clean_chapter_content(item["content"], item["title"])
        soup = BeautifulSoup(_markdown_html(body_markdown), "html.parser")
        for element in soup.find_all(recursive=False):
            if not isinstance(element, Tag):
                continue
            if element.name == "hr":
                flow.append(Paragraph("#", scene_break))
                continue
            markup = _reportlab_inline(element).strip()
            if markup:
                flow.append(Paragraph(markup, body))
    document = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=profile.title,
        author=profile.author_name,
    )
    document.build(
        flow,
        onFirstPage=lambda canvas, doc: _pdf_page_header(canvas, doc, profile=profile),
        onLaterPages=lambda canvas, doc: _pdf_page_header(canvas, doc, profile=profile),
    )


def _plain_document(item: dict) -> str:
    body = _clean_chapter_content(item["content"], item["title"])
    soup = BeautifulSoup(_markdown_html(body), "html.parser")
    text = soup.get_text("\n")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return f"{item['title']}\n\n{text}".strip()


def _trim_words(documents: list[dict], limit: int) -> list[dict]:
    if limit <= 0:
        return []
    remaining = limit
    result: list[dict] = []
    for item in documents:
        words = re.findall(r"\S+", item["content"])
        if len(words) <= remaining:
            result.append(item)
            remaining -= len(words)
        else:
            clipped = " ".join(words[:remaining])
            result.append({**item, "content": clipped})
            remaining = 0
        if remaining <= 0:
            break
    return result


def _sample_documents(
    documents: list[dict],
    destination: SubmissionDestination,
) -> list[dict]:
    if destination.sample_kind == "chapters":
        return documents[: destination.sample_count]
    if destination.sample_kind == "words":
        return _trim_words(documents, destination.sample_count)
    if destination.sample_kind == "full":
        return documents
    if destination.sample_kind == "none":
        return []
    return documents


def _document_material(path: Path, text: str, profile: SubmissionProfile) -> None:
    doc = Document()
    _configure_submission_doc(doc, profile=profile, include_header=False)
    for block in re.split(r"\n\s*\n", text.strip()):
        paragraph = doc.add_paragraph(block.strip())
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(8)
    doc.save(path)


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_page_sample(
    directory: Path,
    *,
    profile: SubmissionProfile,
    documents: list[dict],
    page_count: int,
) -> tuple[list[Path], str]:
    with TemporaryDirectory() as temporary:
        full_pdf = Path(temporary) / "manuscript.pdf"
        _write_manuscript_pdf(full_pdf, profile=profile, documents=documents)
        reader = PdfReader(str(full_pdf))
        selected = reader.pages[: min(page_count, len(reader.pages))]
        pdf_path = directory / f"sample-first-{page_count}-pages.pdf"
        writer = PdfWriter()
        text_pages: list[str] = []
        for page in selected:
            writer.add_page(page)
            text_pages.append(page.extract_text() or "")
        with pdf_path.open("wb") as stream:
            writer.write(stream)
        text = "\n\n".join(part.strip() for part in text_pages if part.strip()).strip()
        txt_path = directory / f"sample-first-{page_count}-pages.txt"
        txt_path.write_text(text + ("\n" if text else ""), encoding="utf-8")
        docx_path = directory / f"sample-first-{page_count}-pages.docx"
        _document_material(docx_path, text, profile)
        return [pdf_path, txt_path, docx_path], text


def _text_digest(content: str) -> str:
    return hashlib.sha256(content.encode("utf-8")).hexdigest()


def _representative_documents(documents: list[dict], limit: int = 40) -> list[dict]:
    if len(documents) <= limit:
        return documents
    if limit <= 1:
        return documents[:1]
    indexes = {
        round(index * (len(documents) - 1) / (limit - 1))
        for index in range(limit)
    }
    return [documents[index] for index in sorted(indexes)]


def _submission_context(slug: str, profile: SubmissionProfile) -> tuple[str, int, bool]:
    documents = compiled_documents(slug)
    by_path = {item["path"]: item for item in documents}
    current_summaries: dict[str, tuple[int, str]] = {}
    try:
        with _connect(slug) as con:
            rows = con.execute(
                "SELECT path, content_hash, summary FROM document_analysis"
            ).fetchall()
        for row in rows:
            item = by_path.get(row["path"])
            summary = str(row["summary"] or "").strip()
            if (
                item is not None
                and summary
                and row["content_hash"] == _text_digest(item["content"])
            ):
                current_summaries[row["path"]] = (chapter_order(row["path"]), summary)
    except sqlite3.Error:
        current_summaries = {}

    lines = [
        f"TITLE: {profile.title}",
        f"GENRE: {profile.genre or '(not specified)'}",
        f"WORD COUNT: {profile.word_count or _word_count(documents):,}",
        f"LOGLINE: {profile.logline or '(not specified)'}",
        f"COMPS: {', '.join(profile.comp_titles) or '(not specified)'}",
        "",
        "STORY CONTEXT IN BINDER ORDER",
    ]
    excerpt_paths = {item["path"] for item in _representative_documents(documents)}
    for item in documents:
        current = current_summaries.get(item["path"])
        if current:
            lines.append(f"\n## {item['title']} [{item['path']}]\nSUMMARY: {current[1]}")
        elif item["path"] in excerpt_paths:
            plain = _plain_document(item)
            lines.append(f"\n## {item['title']} [{item['path']}]\nEXCERPT: {plain[:1800]}")
        else:
            lines.append(
                f"\n## {item['title']} [{item['path']}]\n"
                "No current chapter summary is available; do not invent its events."
            )
    return "\n".join(lines), len(documents), bool(current_summaries)


async def draft_submission_material(
    slug: str,
    *,
    kind: str,
    provider: ProviderConfig,
    profile: SubmissionProfile,
    destination_id: str | None = None,
) -> tuple[str, int, bool]:
    _ensure_project(slug)
    destination = _destination(profile, destination_id) if destination_id else None
    context, document_count, used_summaries = _submission_context(slug, profile)
    target = ""
    if destination:
        target = (
            f"\nDESTINATION: {destination.name} ({destination.kind})\n"
            f"CONTACT: {destination.contact_name}\n"
            f"GUIDELINES NOTES: {destination.notes}\n"
        )
    instructions = {
        "query": (
            "Draft a polished literary query letter. Include a concise hook/story pitch, "
            "title/genre/word count, and bio only from supplied facts. Do not invent awards, "
            "publishing credits, representation history, comp titles, credentials, or "
            "personalization. If personalization is missing, omit it rather than fabricate it."
        ),
        "synopsis": (
            "Draft a clear full-story fiction synopsis that includes major turns, character "
            "causality, climax, and ending/spoilers. Do not hide the ending. Do not invent "
            "connective events that are absent from the supplied story summaries/excerpts."
        ),
        "pitch": (
            "Draft a compact pitch suitable for a query form: protagonist, destabilizing problem, "
            "goal, stakes, central complication, and genre promise. Keep it specific and avoid "
            "empty marketing superlatives."
        ),
        "bio": (
            "Draft a concise author bio using only the author facts supplied below. Do not invent "
            "credentials, awards, occupations, publications, locations, or memberships."
        ),
    }[kind]
    author_facts = (
        f"\nAUTHOR FACTS\nName: {profile.author_name}\n"
        f"Existing bio: {profile.bio or '(none)'}\n"
        f"Website: {profile.website or '(none)'}\n"
    )
    messages = [
        {
            "role": "system",
            "content": (
                "You are EmberWriter's submission editor. Produce professional publishing "
                "material grounded only in the supplied manuscript facts and author facts. "
                "Never fabricate credentials, story facts, agent preferences, or submission "
                "requirements. Return only the requested draft, without commentary."
            ),
        },
        {
            "role": "user",
            "content": f"TASK\n{instructions}{target}{author_facts}\n\nSTORY CONTEXT\n{context}",
        },
    ]
    text = await generate(provider, messages, temperature=0.45, top_p=0.9)
    return text.strip(), document_count, used_summaries


def _follow_up_from_sent(
    destination: SubmissionDestination,
    submitted_at: str,
) -> str:
    if destination.expected_response_days is None:
        return ""
    if submitted_at and _valid_iso_date(submitted_at[:10]):
        sent_date = date.fromisoformat(submitted_at[:10])
    else:
        sent_date = datetime.now(UTC).date()
    return (sent_date + timedelta(days=destination.expected_response_days)).isoformat()


def update_submission_record(
    slug: str,
    profile: SubmissionProfile,
    record_id: str,
    *,
    status: str | None = None,
    submitted_at: str | None = None,
    follow_up_on: str | None = None,
    response_at: str | None = None,
    notes: str | None = None,
) -> SubmissionProfile:
    now = utc_now()
    updated_records: list[SubmissionRecord] = []
    found = False
    for record in profile.records:
        if record.id != record_id:
            updated_records.append(record)
            continue
        found = True
        destination = _destination(profile, record.destination_id)
        changes: dict = {"updated_at": now}
        if status is not None:
            changes["status"] = status
        if submitted_at is not None:
            changes["submitted_at"] = submitted_at
        if follow_up_on is not None:
            changes["follow_up_on"] = follow_up_on
        if response_at is not None:
            changes["response_at"] = response_at
        if notes is not None:
            changes["notes"] = notes
        if status == "sent":
            sent_value = submitted_at if submitted_at is not None else record.submitted_at
            if not sent_value:
                sent_value = datetime.now(UTC).date().isoformat()
                changes["submitted_at"] = sent_value
            if follow_up_on is None and not record.follow_up_on:
                changes["follow_up_on"] = _follow_up_from_sent(destination, sent_value)
        updated_records.append(record.model_copy(update=changes))
    if not found:
        raise FileNotFoundError(record_id)
    saved = profile.model_copy(update={"records": updated_records})
    return save_submission_profile(slug, saved)


def build_submission_package(
    slug: str,
    *,
    profile: SubmissionProfile,
    destination_id: str,
) -> tuple[SubmissionBuildResponse, SubmissionProfile]:
    validation = validate_submission(slug, profile, destination_id)
    if not validation.valid:
        errors = "; ".join(
            issue.message for issue in validation.issues if issue.level == "error"
        )
        raise ValueError(errors or "Submission profile is not ready")
    destination = _destination(profile, destination_id)
    documents = compiled_documents(slug)
    package_id = (
        datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        + "-submission-"
        + uuid4().hex[:8]
    )
    root = project_root(slug)
    directory = root / "exports" / package_id
    directory.mkdir(parents=True, exist_ok=True)
    included: list[Path] = []

    destination_path = directory / "destination.json"
    destination_path.write_text(
        json.dumps(destination.model_dump(), indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    included.append(destination_path)

    if profile.query_letter.strip():
        query_txt = directory / "query-letter.txt"
        query_docx = directory / "query-letter.docx"
        query_txt.write_text(profile.query_letter.strip() + "\n", encoding="utf-8")
        _document_material(query_docx, profile.query_letter, profile)
        included.extend([query_txt, query_docx])
    if profile.synopsis.strip():
        synopsis_txt = directory / "synopsis.txt"
        synopsis_docx = directory / "synopsis.docx"
        synopsis_txt.write_text(profile.synopsis.strip() + "\n", encoding="utf-8")
        _document_material(synopsis_docx, profile.synopsis, profile)
        included.extend([synopsis_txt, synopsis_docx])
    if profile.bio.strip():
        bio_txt = directory / "author-bio.txt"
        bio_txt.write_text(profile.bio.strip() + "\n", encoding="utf-8")
        included.append(bio_txt)

    sample_text = ""
    if destination.sample_kind == "pages":
        page_files, sample_text = _write_page_sample(
            directory,
            profile=profile,
            documents=documents,
            page_count=destination.sample_count,
        )
        included.extend(page_files)
    elif destination.sample_kind != "none":
        sample_documents = _sample_documents(documents, destination)
        sample_docx = directory / "manuscript-sample.docx"
        sample_pdf = directory / "manuscript-sample.pdf"
        _write_manuscript_docx(
            sample_docx,
            profile=profile,
            documents=sample_documents,
            include_title_page=destination.sample_kind == "full",
        )
        _write_manuscript_pdf(sample_pdf, profile=profile, documents=sample_documents)
        sample_text = "\n\n\n".join(_plain_document(item) for item in sample_documents)
        sample_txt = directory / "manuscript-sample.txt"
        sample_txt.write_text(sample_text + ("\n" if sample_text else ""), encoding="utf-8")
        included.extend([sample_docx, sample_pdf, sample_txt])

    if destination.attachment_mode in {"body", "mixed"}:
        body_parts = []
        if destination.query_required and profile.query_letter.strip():
            body_parts.append(profile.query_letter.strip())
        if destination.synopsis_required and profile.synopsis.strip():
            body_parts.append("SYNOPSIS\n\n" + profile.synopsis.strip())
        if sample_text.strip():
            body_parts.append(sample_text.strip())
        email_body = directory / "submission-body.txt"
        email_body.write_text("\n\n---\n\n".join(body_parts) + "\n", encoding="utf-8")
        included.append(email_body)

    manifest_entries = []
    for path in included:
        manifest_entries.append(
            {
                "path": path.name,
                "bytes": path.stat().st_size,
                "sha256": _hash(path),
            }
        )
    manifest = {
        "schema_version": 1,
        "package_id": package_id,
        "generated_at": datetime.now(UTC).isoformat(),
        "rules_reviewed": RULES_REVIEWED,
        "destination": destination.model_dump(),
        "manuscript_preset": profile.manuscript_preset,
        "preset_reference": SHUNN_FORMAT_URL,
        "files": manifest_entries,
        "warning": (
            "Destination-specific current guidelines override Ember presets. "
            "Verify every file and field before sending."
        ),
    }
    manifest_path = directory / "submission-manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    included.append(manifest_path)

    readme = directory / "README.txt"
    readme.write_text(
        "EmberWriter traditional submission handoff\n\n"
        f"Destination: {destination.name}\n"
        f"Package: {package_id}\n"
        f"Guidelines: {destination.guidelines_url or '(not recorded)'}\n"
        f"Sample: {validation.sample_description}\n\n"
        "This package is a preparation aid, not an automated submission. Re-open the files "
        "and compare them to the destination's current guidelines before sending.\n",
        encoding="utf-8",
    )
    included.append(readme)

    zip_path = directory / "submission-package.zip"
    with zipfile.ZipFile(
        zip_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as archive:
        for path in included:
            archive.write(path, path.name)

    now = utc_now()
    record = SubmissionRecord(
        id=uuid4().hex,
        destination_id=destination.id,
        status="ready",
        package_id=package_id,
        follow_up_on="",
        created_at=now,
        updated_at=now,
    )
    updated_profile = profile.model_copy(update={"records": [*profile.records, record]})
    save_submission_profile(slug, updated_profile)

    artifacts = [
        SubmissionPackageArtifact(
            format="zip",
            filename=zip_path.name,
            relative_path=str(zip_path.relative_to(root)).replace("\\", "/"),
            bytes=zip_path.stat().st_size,
        ),
        SubmissionPackageArtifact(
            format="json",
            filename=manifest_path.name,
            relative_path=str(manifest_path.relative_to(root)).replace("\\", "/"),
            bytes=manifest_path.stat().st_size,
        ),
    ]
    for path in included:
        if path.suffix.casefold() not in {".docx", ".pdf", ".txt"}:
            continue
        artifacts.append(
            SubmissionPackageArtifact(
                format=path.suffix.casefold().lstrip("."),
                filename=path.name,
                relative_path=str(path.relative_to(root)).replace("\\", "/"),
                bytes=path.stat().st_size,
            )
        )
    response = SubmissionBuildResponse(
        package_id=package_id,
        validation=validation,
        artifacts=artifacts,
        included_files=[path.name for path in included],
    )
    return response, updated_profile
