from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from ebooklib import epub
from pypdf import PdfReader

from app import binder, ingest, publishing, revisions, storage
from app.binder_models import BinderNodeCreate, BinderNodeUpdate, BinderReorderRequest


def use_temp_data(tmp_path: Path) -> None:
    storage.DATA_ROOT = tmp_path
    storage.PROJECTS_ROOT = tmp_path / "projects"


def test_revision_history_diff_and_restore_preserves_later_history(tmp_path: Path) -> None:
    use_temp_data(tmp_path)
    project = storage.create_project("Revision Test")
    slug = project["slug"]
    path = "manuscript/chapter-001.md"

    first_text = "# Chapter 1\n\nThe old opening remains recoverable.\n"
    second_text = "# Chapter 1\n\nThe new opening changes everything.\n"
    storage.save_text(slug, path, first_text)
    first = revisions.record_revision(slug, path, first_text, source="save")
    storage.save_text(slug, path, second_text)
    second = revisions.record_revision(slug, path, second_text, source="save")

    history = revisions.list_revisions(slug, path)
    assert [item["id"] for item in history[:2]] == [second["id"], first["id"]]

    diff = revisions.compare_revisions(slug, first["id"], second["id"])["diff"]
    assert "-The old opening remains recoverable." in diff
    assert "+The new opening changes everything." in diff

    restored = revisions.restore_revision(slug, first["id"], "Undo experiment")
    assert storage.read_text(slug, path) == first_text
    assert restored["revision"]["source"] == "restore"

    ids_after_restore = [item["id"] for item in revisions.list_revisions(slug, path)]
    assert first["id"] in ids_after_restore
    assert second["id"] in ids_after_restore
    assert restored["revision"]["id"] == ids_after_restore[0]


def test_docx_novel_import_splits_chapters_and_creates_baseline_revisions(tmp_path: Path) -> None:
    use_temp_data(tmp_path)
    project = storage.create_project("Import Test")
    slug = project["slug"]

    document = Document()
    document.add_heading("Chapter 1", level=1)
    document.add_paragraph("Mara enters the winter city.")
    document.add_heading("Chapter 2", level=1)
    document.add_paragraph("At dawn, the bargain comes due.")
    buffer = io.BytesIO()
    document.save(buffer)

    result = ingest.import_bytes(slug, "draft.docx", buffer.getvalue(), "novel")
    assert len(result["documents"]) == 2
    assert result["documents"][0]["title"] == "Chapter 1"
    assert result["documents"][1]["title"] == "Chapter 2"

    for imported in result["documents"]:
        assert imported["path"].startswith("manuscript/")
        history = revisions.list_revisions(slug, imported["path"])
        assert len(history) == 1
        assert history[0]["source"] == "import"

    state = binder.get_binder(slug)
    imported_ids = {item["node_id"] for item in result["documents"]}
    assert imported_ids.issubset({node.id for node in state.nodes})


def test_pasted_ideas_and_research_land_outside_compile_draft(tmp_path: Path) -> None:
    use_temp_data(tmp_path)
    project = storage.create_project("Idea Test")
    slug = project["slug"]

    idea = ingest.import_pasted_text(
        slug,
        "Ending possibility",
        "The crown is never destroyed.",
        "idea",
    )
    research = ingest.import_pasted_text(
        slug,
        "Mars Notes",
        "Orbital periods and dust storms.",
        "research",
    )
    state = binder.get_binder(slug)
    nodes = {node.id: node for node in state.nodes}

    idea_node = nodes[idea["documents"][0]["node_id"]]
    research_node = nodes[research["documents"][0]["node_id"]]
    assert idea_node.include_in_compile is False
    assert research_node.include_in_compile is False
    assert idea_node.path and idea_node.path.startswith("notes/")
    assert research_node.path and research_node.path.startswith("research/")


def test_compile_order_drives_docx_epub_and_pdf_exports(tmp_path: Path) -> None:
    use_temp_data(tmp_path)
    project = storage.create_project("Publish Test")
    slug = project["slug"]
    state = binder.get_binder(slug)
    draft = next(node for node in state.nodes if node.title == "Draft")
    first = next(node for node in state.nodes if node.path == "manuscript/chapter-001.md")
    assert first.path
    storage.save_text(slug, first.path, "# First Scene\n\nThe first scene text.\n")
    state = binder.update_node(slug, first.id, BinderNodeUpdate(title="First Scene"))

    before = {node.id for node in state.nodes}
    state = binder.create_node(
        slug,
        BinderNodeCreate(title="Second Scene", kind="document", parent_id=draft.id),
    )
    second = next(node for node in state.nodes if node.id not in before)
    assert second.path
    storage.save_text(slug, second.path, "# Second Scene\n\nThe second scene text.\n")

    siblings = sorted(
        (node for node in state.nodes if node.parent_id == draft.id),
        key=lambda node: node.position,
    )
    binder.reorder_nodes(
        slug,
        BinderReorderRequest(
            parent_id=draft.id,
            node_ids=[second.id, *[node.id for node in siblings if node.id != second.id]],
        ),
    )
    ordered = publishing.compiled_documents(slug)
    assert ordered[0]["title"] == "Second Scene"

    result = publishing.export_project(
        slug,
        formats=["docx", "epub", "pdf"],
        title="The Ember Test",
        author="A. Writer",
        language="en",
        include_toc=True,
        trim_width=6,
        trim_height=9,
    )
    root = storage.project_root(slug)
    paths = {
        artifact["format"]: root / artifact["relative_path"]
        for artifact in result["artifacts"]
    }

    assert paths["docx"].read_bytes()[:2] == b"PK"
    assert paths["epub"].read_bytes()[:2] == b"PK"
    assert paths["pdf"].read_bytes()[:4] == b"%PDF"

    exported_docx = Document(paths["docx"])
    docx_text = "\n".join(paragraph.text for paragraph in exported_docx.paragraphs)
    assert docx_text.index("Second Scene") < docx_text.index("First Scene")

    exported_epub = epub.read_epub(str(paths["epub"]))
    epub_titles = [
        item.title
        for item in exported_epub.get_items()
        if getattr(item, "title", None)
    ]
    assert "Second Scene" in epub_titles

    pdf = PdfReader(str(paths["pdf"]))
    pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert pdf_text.index("Second Scene") < pdf_text.index("First Scene")
